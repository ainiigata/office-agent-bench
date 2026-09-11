"""excel-01 のフィクスチャ生成。python3 make_fixtures.py で input/ と expected/ を再生成する。"""
from __future__ import annotations

import datetime
import json
import random
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
FILE = "売上明細_2026上期.xlsx"
STAFF = ["佐藤 健一", "鈴木 美咲", "高橋 大輔", "田中 陽子", "伊藤 翔太"]
CLIENTS = ["株式会社アルファ商事", "ベータ工業株式会社", "ガンマ物産株式会社", "有限会社デルタ企画",
           "イプシロン建設株式会社", "株式会社ゼータ食品", "イータ電機株式会社", "シータ運輸株式会社"]
PRODUCTS = [("業務用コピー用紙(A4・5箱)", 4200), ("トナーカートリッジ", 12800), ("オフィスチェア", 24500),
            ("スチール書庫", 38000), ("会議用テーブル", 56000), ("LEDデスクライト", 6800)]
MONTHS = list(range(4, 10))


def main() -> None:
    rnd = random.Random(20260911)
    rows = []
    for _ in range(300):
        day = datetime.date(2026, 4, 1) + datetime.timedelta(days=rnd.randrange(183))  # 4/1〜9/30
        product, price = rnd.choice(PRODUCTS)
        qty = rnd.randint(1, 30)
        rows.append([day, rnd.choice(STAFF), rnd.choice(CLIENTS), product, qty, price, qty * price])
    rows.sort(key=lambda r: r[0])

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "明細"
    ws.append(["日付", "担当者", "得意先", "商品", "数量", "単価", "金額"])
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append(r)
    for (cell,) in ws.iter_rows(min_row=2, min_col=1, max_col=1):
        cell.number_format = "yyyy/mm/dd"
    for i, w in enumerate([12, 12, 26, 28, 6, 10, 12], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    (HERE / "input").mkdir(exist_ok=True)
    (HERE / "expected").mkdir(exist_ok=True)
    wb.save(HERE / "input" / FILE)

    pivot = {s: {m: 0 for m in MONTHS} for s in STAFF}
    for day, staff, *_rest, amount in rows:
        pivot[staff][day.month] += amount
    grand = sum(sum(v.values()) for v in pivot.values())

    ws2 = wb.create_sheet("集計")
    ws2.append(["担当者"] + [f"{m}月" for m in MONTHS] + ["合計", "構成比"])
    for s in STAFF:
        total = sum(pivot[s].values())
        ws2.append([s] + [pivot[s][m] for m in MONTHS] + [total, round(total / grand, 4)])
    ws2.append(["合計"] + [sum(pivot[s][m] for s in STAFF) for m in MONTHS] + [grand, 1.0])
    for c in ws2[1]:
        c.font = Font(bold=True)
    for row in ws2.iter_rows(min_row=2, min_col=2, max_col=8):
        for c in row:
            c.number_format = "#,##0"
    for (c,) in ws2.iter_rows(min_row=2, min_col=9, max_col=9):
        c.number_format = "0.0%"
    ws2.column_dimensions["A"].width = 12
    wb.save(HERE / "expected" / FILE)

    # ダイジェストは check.py と同じ関数で、保存後のブックから計算する(日付の型を揃えるため)
    from check import detail_digest

    saved = openpyxl.load_workbook(HERE / "expected" / FILE, data_only=True)["明細"]
    values = {"staff": STAFF, "months": MONTHS,
              "pivot": {s: {str(m): pivot[s][m] for m in MONTHS} for s in STAFF},
              "grand_total": grand, "detail_rows": len(rows), "detail_digest": detail_digest(saved)}
    (HERE / "expected" / "values.json").write_text(json.dumps(values, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"生成: {len(rows)}行, 総額 {grand:,}")


if __name__ == "__main__":
    main()
