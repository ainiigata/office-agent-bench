"""update-01 のフィクスチャ生成。python3 make_fixtures.py で input/ と expected/ を再生成する。"""
from __future__ import annotations

import json
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, Side

HERE = Path(__file__).resolve().parent
FILE_IN = "令和7年度_委員名簿.xlsx"
FILE_OUT = "令和8年度_委員名簿.xlsx"
TERM = "令和7年4月1日〜令和9年3月31日"
TERM_NEW = "令和8年4月1日〜令和9年3月31日"
HEADER = ["No", "区分", "氏名", "所属", "役職", "電話", "任期"]

MEMBERS_R7 = [
    ["学識経験者", "渡辺 隆", "みなと臨海大学", "教授", "025-111-1001", TERM],
    ["学識経験者", "中島 京子", "青波学園大学", "准教授", "025-111-1002", TERM],
    ["関係団体", "佐々木 一郎", "みなと市医師会", "理事", "025-222-2001", TERM],
    ["関係団体", "木村 幸子", "みなと市社会福祉協議会", "事務局次長", "025-222-2002", TERM],
    ["関係団体", "加藤 由美", "みなと市ボランティア協会", "会長", "025-222-2003", TERM],
    ["関係団体", "松本 誠治", "みなと市民生委員児童委員協議会", "会長", "025-222-2004", TERM],
    ["関係団体", "井上 和也", "みなと市老人クラブ連合会", "副会長", "025-222-2005", TERM],
    ["関係団体", "山本 千尋", "みなと市障がい者団体連絡協議会", "事務局長", "025-222-2006", TERM],
    ["公募委員", "小林 正雄", "北地区在住", "", "", TERM],
    ["公募委員", "斎藤 美香", "東地区在住", "", "", TERM],
    ["公募委員", "森 拓也", "西地区在住", "", "", TERM],
    ["行政", "山口 誠", "みなと市役所", "健康福祉課長", "025-333-3001", TERM],
    ["行政", "清水 里奈", "福祉部", "地域福祉課長", "025-333-3002", TERM],
    ["行政", "阿部 健太", "こども未来部", "こども政策課長", "025-333-3003", TERM],
    ["行政", "石川 直子", "保健衛生部", "健康増進課長", "025-333-3004", TERM],
]

MEMBERS_R8 = [
    ["学識経験者", "渡辺 隆", "みなと臨海大学", "教授", "025-111-1099", TERM],
    ["学識経験者", "中島 京子", "青波学園大学", "准教授", "025-111-1002", TERM],
    ["関係団体", "中村 恵子", "みなと市医師会", "理事", "025-222-2001", TERM_NEW],
    ["関係団体", "林 大樹", "みなと市社会福祉協議会", "事務局次長", "025-222-2002", TERM_NEW],
    ["関係団体", "加藤 由美", "みなと市ボランティア・市民活動センター", "会長", "025-222-2003", TERM],
    ["関係団体", "松本 誠治", "みなと市民生委員児童委員協議会", "会長", "025-222-2004", TERM],
    ["関係団体", "井上 和也", "みなと市老人クラブ連合会", "副会長", "025-222-2005", TERM],
    ["関係団体", "山本 千尋", "みなと市障がい者団体連絡協議会", "事務局長", "025-222-2006", TERM],
    ["公募委員", "斎藤 美香", "東地区在住", "", "", TERM],
    ["公募委員", "森 拓也", "西地区在住", "", "", TERM],
    ["行政", "岡田 浩二", "みなと市役所", "健康福祉課長", "025-333-3001", TERM_NEW],
    ["行政", "清水 里奈", "福祉部", "地域福祉課長", "025-333-3002", TERM],
    ["行政", "阿部 健太", "こども未来部", "こども政策課長", "025-333-3003", TERM],
    ["行政", "石川 直子", "保健衛生部", "健康増進課長", "025-333-3004", TERM],
]

NOTICE = """地域福祉推進協議会 事務局 各位

令和8年度の委員名簿の更新について、以下のとおり異動の連絡がありましたので反映をお願いします。
(いずれも令和8年4月1日付)

1. 委員の交代
 ・みなと市医師会: 佐々木 一郎 委員(理事)が退任し、後任は 中村 恵子 氏(理事)
 ・みなと市社会福祉協議会: 木村 幸子 委員(事務局次長)が退任し、後任は 林 大樹 氏(事務局次長)
 ・みなと市役所: 山口 誠 委員(健康福祉課長)が人事異動により退任し、後任は 岡田 浩二 氏(健康福祉課長)
 ※後任の方の任期は、いずれも前任者の残任期間(令和8年4月1日〜令和9年3月31日)です。
 ※電話番号は前任者と同じ(所属の代表番号)です。

2. 退任(後任なし)
 ・公募委員 小林 正雄 委員(一身上の都合により退任)

3. 電話番号の変更
 ・渡辺 隆 委員: 025-111-1001 → 025-111-1099

4. 所属名の変更
 ・加藤 由美 委員の所属「みなと市ボランティア協会」は、令和8年4月1日から「みなと市ボランティア・市民活動センター」に名称変更
"""

VALUES = {
    "file_in": FILE_IN, "file_out": FILE_OUT,
    "retired": ["小林 正雄", "佐々木 一郎", "木村 幸子", "山口 誠"],
    "new_members": [
        {"name": "中村 恵子", "org": "みなと市医師会", "title": "理事", "term_from": "令和8年4月1日", "term_to": "令和9年3月31日"},
        {"name": "林 大樹", "org": "みなと市社会福祉協議会", "title": "事務局次長", "term_from": "令和8年4月1日", "term_to": "令和9年3月31日"},
        {"name": "岡田 浩二", "org": "みなと市役所", "title": "健康福祉課長", "term_from": "令和8年4月1日", "term_to": "令和9年3月31日"},
    ],
    "phone_change": {"name": "渡辺 隆", "old": "025-111-1001", "new": "025-111-1099"},
    "org_change": {"name": "加藤 由美", "old": "みなと市ボランティア協会", "new": "みなと市ボランティア・市民活動センター"},
    "member_count": len(MEMBERS_R8),
}


def build(path: Path, year: str, members: list[list[str]]) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "委員名簿"
    ws["A1"] = f"{year} 地域福祉推進協議会 委員名簿"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"({year[:-1]}4月1日現在)"  # 例: (令和7年4月1日現在)
    ws.append([])
    ws.append(HEADER)
    thin = Side(style="thin")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for c in ws[4]:
        c.font = Font(bold=True)
        c.border = border
        c.alignment = Alignment(horizontal="center")
    for i, m in enumerate(members, start=1):
        ws.append([i] + m)
        for c in ws[ws.max_row]:
            c.border = border
    for col, w in zip("ABCDEFG", [5, 12, 14, 34, 16, 14, 30]):
        ws.column_dimensions[col].width = w
    wb.save(path)


def main() -> None:
    (HERE / "input").mkdir(exist_ok=True)
    (HERE / "expected").mkdir(exist_ok=True)
    build(HERE / "input" / FILE_IN, "令和7年度", MEMBERS_R7)
    (HERE / "input" / "異動連絡.txt").write_text(NOTICE, encoding="utf-8")
    build(HERE / "expected" / FILE_OUT, "令和8年度", MEMBERS_R8)
    (HERE / "expected" / "values.json").write_text(json.dumps(VALUES, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"生成完了: R7 {len(MEMBERS_R7)}名 → R8 {len(MEMBERS_R8)}名")


if __name__ == "__main__":
    main()
