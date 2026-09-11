"""web-02 のフィクスチャ生成。python3 make_fixtures.py で expected/ を再生成する。

正解は内閣府の公開CSV(https://www8.cao.go.jp/chosei/shukujitsu/syukujitsu.csv)で確認した値。
"""
from __future__ import annotations

import datetime
import json
from pathlib import Path

import openpyxl
from openpyxl.styles import Font

HERE = Path(__file__).resolve().parent
VERIFIED_ON = "2026-09-11"
SOURCE = "https://www8.cao.go.jp/chosei/shukujitsu/syukujitsu.csv"
WEEKDAYS = "月火水木金土日"
HOLIDAYS = [
    ("2027-01-01", "元日"), ("2027-01-11", "成人の日"), ("2027-02-11", "建国記念の日"),
    ("2027-02-23", "天皇誕生日"), ("2027-03-21", "春分の日"), ("2027-03-22", "休日"),
    ("2027-04-29", "昭和の日"), ("2027-05-03", "憲法記念日"), ("2027-05-04", "みどりの日"),
    ("2027-05-05", "こどもの日"), ("2027-07-19", "海の日"), ("2027-08-11", "山の日"),
    ("2027-09-20", "敬老の日"), ("2027-09-23", "秋分の日"), ("2027-10-11", "スポーツの日"),
    ("2027-11-03", "文化の日"), ("2027-11-23", "勤労感謝の日"),
]


def main() -> None:
    (HERE / "expected").mkdir(exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "祝日2027"
    ws.append(["日付", "曜日", "名称", "備考"])
    for c in ws[1]:
        c.font = Font(bold=True)
    holidays = []
    for iso, name in HOLIDAYS:
        d = datetime.date.fromisoformat(iso)
        substitute = name == "休日"
        note = "振替休日" if substitute else ("土曜" if d.weekday() == 5 else "")
        ws.append([d, WEEKDAYS[d.weekday()], "振替休日" if substitute else name, note])
        ws.cell(row=ws.max_row, column=1).number_format = "yyyy-mm-dd"
        holidays.append({"date": iso, "name": name, "substitute": substitute})
    ws.append([])
    ws.append([f"出典: {SOURCE}"])
    ws.append([f"確認日: {VERIFIED_ON}"])
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["C"].width = 16
    ws.column_dimensions["D"].width = 12
    wb.save(HERE / "expected" / "祝日2027.xlsx")
    values = {"verified_on": VERIFIED_ON, "source": SOURCE, "holidays": holidays}
    (HERE / "expected" / "values.json").write_text(json.dumps(values, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"生成完了: {len(holidays)}行")


if __name__ == "__main__":
    main()
