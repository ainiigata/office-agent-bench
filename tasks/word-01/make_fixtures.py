"""word-01 のフィクスチャ生成。python3 make_fixtures.py で input/ と expected/ を再生成する。"""
from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

HERE = Path(__file__).resolve().parent

MEMO = """【防災訓練の案内通知 作成メモ】
文書番号: 中総第123号
日付: 令和8年9月15日
宛名: 各自治会長 様
発信者: みなと市役所 総務課長
件名: 令和8年度 地域防災訓練の実施について(案内)
日時: 令和8年10月18日(日) 午前9時から正午まで
場所: 中央小学校グラウンド
対象: 各自治会の会長および防災担当者(自治会ごとに2名まで)
内容: 初期消火訓練、避難誘導訓練、炊き出し体験
持ち物: 動きやすい服装、飲み物、筆記用具
雨天時: 小雨決行。荒天の場合は前日17時までに各自治会長へ電話で連絡する
回答: 参加人数を10月9日(金)までにメールで返信してほしい
問い合わせ先: みなと市役所 総務課 防災係(電話 025-123-4567)
"""

STYLE = """【公文書(通知文)の基本レイアウト例】

                                            ○○第○○号
                                            令和○年○月○日
○○○○ 様
                                            ○○市役所 ○○課長

              ○○○○について(通知)

  日頃から○○にご協力いただき、厚くお礼申し上げます。
  さて、○○を下記のとおり実施しますので、ご参加くださるようお願いします。

                        記

1 日時  ○○
2 場所  ○○
3 対象  ○○
4 持ち物 ○○
5 その他 ○○

                                                            以上

【担当】○○課○○係 電話 ○○○-○○○-○○○○
"""

VALUES = {
    "header": {"文書番号": "中総第123号", "日付": "令和8年9月15日", "宛名": "各自治会長",
               "発信者": "総務課長", "件名": "地域防災訓練の実施について"},
    "body": {"日時": "令和8年10月18日", "場所": "中央小学校グラウンド", "持ち物": "筆記用具"},
    "phone_digits": "0251234567",
}


def _p(doc, text, align=None, bold=False, size=None):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    if size:
        run.font.size = Pt(size)
    if align:
        p.alignment = align
    return p


def build_expected(path: Path) -> None:
    doc = Document()
    _p(doc, "中総第123号", WD_ALIGN_PARAGRAPH.RIGHT)
    _p(doc, "令和8年9月15日", WD_ALIGN_PARAGRAPH.RIGHT)
    _p(doc, "各自治会長 様")
    _p(doc, "みなと市役所 総務課長", WD_ALIGN_PARAGRAPH.RIGHT)
    _p(doc, "")
    _p(doc, "令和8年度 地域防災訓練の実施について(案内)", WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=13)
    _p(doc, "")
    _p(doc, "　日頃から地域防災活動にご協力いただき、厚くお礼申し上げます。")
    _p(doc, "　さて、令和8年度地域防災訓練を下記のとおり実施しますので、ご参加くださるようお願いします。"
            "つきましては、参加人数を10月9日(金)までにメールでご回答ください。")
    _p(doc, "")
    _p(doc, "記", WD_ALIGN_PARAGRAPH.CENTER)
    _p(doc, "")
    _p(doc, "1 日時　令和8年10月18日(日) 午前9時から正午まで")
    _p(doc, "2 場所　中央小学校グラウンド")
    _p(doc, "3 対象　各自治会の会長および防災担当者(自治会ごとに2名まで)")
    _p(doc, "4 内容　初期消火訓練、避難誘導訓練、炊き出し体験")
    _p(doc, "5 持ち物　動きやすい服装、飲み物、筆記用具")
    _p(doc, "6 雨天時　小雨決行。荒天の場合は前日17時までに各自治会長へ電話でご連絡します。")
    _p(doc, "")
    _p(doc, "以上", WD_ALIGN_PARAGRAPH.RIGHT)
    _p(doc, "")
    _p(doc, "【担当】みなと市役所 総務課 防災係　電話 025-123-4567")
    doc.save(path)


def main() -> None:
    (HERE / "input").mkdir(exist_ok=True)
    (HERE / "expected").mkdir(exist_ok=True)
    (HERE / "input" / "メモ.txt").write_text(MEMO, encoding="utf-8")
    (HERE / "input" / "参考_様式例.txt").write_text(STYLE, encoding="utf-8")
    build_expected(HERE / "expected" / "通知文.docx")
    (HERE / "expected" / "values.json").write_text(json.dumps(VALUES, ensure_ascii=False, indent=2), encoding="utf-8")
    print("生成完了")


if __name__ == "__main__":
    main()
