import pytest
from docx import Document

from bench.docx_text import docx_text


def test_docx_text_joins_paragraphs_and_tables(tmp_path):
    doc = Document()
    doc.add_paragraph("一行目")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "セルA"
    table.cell(0, 1).text = "セルB"
    doc.add_paragraph("三行目")
    p = tmp_path / "t.docx"
    doc.save(p)
    text = docx_text(p)
    assert "一行目" in text and "セルA" in text and "セルB" in text and "三行目" in text
    assert text.index("一行目") < text.index("セルA") < text.index("三行目")


def test_docx_text_missing_raises(tmp_path):
    with pytest.raises(FileNotFoundError, match="成果物がありません"):
        docx_text(tmp_path / "no.docx")
