"""実行キット各タスクのチェッカーを、模範成果物・未着手・壊した成果物で検証する。"""
from pathlib import Path

import pytest

from tools.score import run_checker
from tools.selftest import build_run_dir

ROOT = Path(__file__).resolve().parents[1]
KIT_TASKS = ["excel-01", "word-01", "mail-03", "web-02", "update-01"]


def _run(task_id: str, tmp_path: Path, with_expected: bool):
    run_dir = build_run_dir(ROOT / "tasks" / task_id, tmp_path / "2000-01-01-test" / task_id, with_expected)
    return run_dir, {i["id"]: i for i in run_checker(ROOT, task_id, run_dir)}


@pytest.mark.parametrize("task_id", KIT_TASKS)
def test_expected_output_passes_everything(task_id, tmp_path):
    _, items = _run(task_id, tmp_path, with_expected=True)
    failed = [i for i in items.values() if not i["pass"]]
    assert not failed, failed


@pytest.mark.parametrize("task_id", KIT_TASKS)
def test_untouched_input_does_not_pass_everything(task_id, tmp_path):
    _, items = _run(task_id, tmp_path, with_expected=False)
    assert any(not i["pass"] for i in items.values())


def test_excel01_detects_wrong_value_and_modified_detail(tmp_path):
    import openpyxl
    run_dir, _ = _run("excel-01", tmp_path, with_expected=True)
    path = run_dir / "売上明細_2026上期.xlsx"
    wb = openpyxl.load_workbook(path)
    ws = wb["集計"]
    ws.cell(row=2, column=2).value = ws.cell(row=2, column=2).value + 5000
    wb["明細"].delete_rows(2)
    wb.save(path)
    items = {i["id"]: i for i in run_checker(ROOT, "excel-01", run_dir)}
    assert items["sheet_exists"]["pass"]
    assert not items["pivot_match"]["pass"]
    assert not items["detail_intact"]["pass"]


def test_excel01_detects_gutted_detail_columns(tmp_path):
    """行数と金額合計が同じでも、明細のセル値を書き換えたら detail_intact は落ちる。"""
    import openpyxl
    run_dir, _ = _run("excel-01", tmp_path, with_expected=True)
    path = run_dir / "売上明細_2026上期.xlsx"
    wb = openpyxl.load_workbook(path)
    ws = wb["明細"]
    for row in range(2, ws.max_row + 1):
        ws.cell(row=row, column=2).value = "担当者不明"  # 担当者を全行つぶす
        ws.cell(row=row, column=3).value = None         # 得意先を空にする
    wb.save(path)
    items = {i["id"]: i for i in run_checker(ROOT, "excel-01", run_dir)}
    assert not items["detail_intact"]["pass"], items["detail_intact"]
    assert "300" in items["detail_intact"]["detail"]  # 行数・金額合計の診断は残す


def test_excel01_formula_without_cached_value_is_reported(tmp_path):
    import openpyxl
    run_dir, _ = _run("excel-01", tmp_path, with_expected=True)
    path = run_dir / "売上明細_2026上期.xlsx"
    wb = openpyxl.load_workbook(path)
    wb["集計"].cell(row=2, column=2).value = "=SUM(1,2)"
    wb.save(path)
    items = {i["id"]: i for i in run_checker(ROOT, "excel-01", run_dir)}
    assert not items["pivot_match"]["pass"]
    assert "数式" in items["pivot_match"]["detail"]


def test_excel01_accepts_text_formatted_numbers(tmp_path):
    """桁区切り付きの文字列・「18.8%」形式の文字列でも数値として読む。"""
    import openpyxl
    run_dir, _ = _run("excel-01", tmp_path, with_expected=True)
    path = run_dir / "売上明細_2026上期.xlsx"
    wb = openpyxl.load_workbook(path)
    ws = wb["集計"]
    for row in range(2, 7):                      # 担当者5名の行
        for col in range(2, 9):                  # 4月〜9月 + 合計
            ws.cell(row=row, column=col).value = f"{ws.cell(row=row, column=col).value:,}円"
        share = ws.cell(row=row, column=9)
        share.value = f"{share.value * 100:.1f}%"
    wb.save(path)
    items = {i["id"]: i for i in run_checker(ROOT, "excel-01", run_dir)}
    assert items["pivot_match"]["pass"], items["pivot_match"]
    assert items["total_match"]["pass"], items["total_match"]
    assert items["share_sum"]["pass"], items["share_sum"]


def test_excel01_accepts_fullwidth_and_gatsudo_headers(tmp_path):
    import openpyxl
    run_dir, _ = _run("excel-01", tmp_path, with_expected=True)
    path = run_dir / "売上明細_2026上期.xlsx"
    wb = openpyxl.load_workbook(path)
    ws = wb["集計"]
    for col, label in zip(range(2, 8), ["４月度", "５月度", "６月度", "７月度", "８月度", "９月度"]):
        ws.cell(row=1, column=col).value = label
    wb.save(path)
    items = {i["id"]: i for i in run_checker(ROOT, "excel-01", run_dir)}
    assert items["pivot_match"]["pass"], items["pivot_match"]
    assert items["total_match"]["pass"], items["total_match"]
    assert items["share_sum"]["pass"], items["share_sum"]


def test_word01_detects_missing_ki_and_header(tmp_path):
    from docx import Document
    run_dir, _ = _run("word-01", tmp_path, with_expected=True)
    doc = Document()
    doc.add_paragraph("令和8年9月15日")
    doc.add_paragraph("各自治会長 様")
    doc.add_paragraph("地域防災訓練を実施します。日時は令和8年10月18日、場所は中央小学校グラウンド、持ち物は筆記用具。")
    doc.add_paragraph("以上")
    doc.add_paragraph("記")
    doc.add_paragraph("問い合わせ 025(123)4567")
    doc.save(run_dir / "通知文.docx")
    items = {i["id"]: i for i in run_checker(ROOT, "word-01", run_dir)}
    assert items["docx_exists"]["pass"]
    assert not items["header_fields"]["pass"] and "中総第123号" in items["header_fields"]["detail"]
    assert not items["ki_ijo_order"]["pass"]
    assert items["body_content"]["pass"]
    assert items["phone"]["pass"]  # 括弧書きの電話番号も数字列で一致とみなす


def test_word01_missing_subject_line_fails_header_fields(tmp_path):
    from docx import Document
    run_dir, _ = _run("word-01", tmp_path, with_expected=True)
    path = run_dir / "通知文.docx"
    doc = Document(str(path))
    for p in doc.paragraphs:
        if "地域防災訓練の実施について" in p.text:
            p._element.getparent().remove(p._element)
    doc.save(path)
    items = {i["id"]: i for i in run_checker(ROOT, "word-01", run_dir)}
    assert not items["header_fields"]["pass"] and "件名" in items["header_fields"]["detail"]
    assert items["body_content"]["pass"]
    assert items["ki_ijo_order"]["pass"]


def test_word01_detects_memo_dump(tmp_path):
    """メモ.txt をそのまま貼って記・以上を足しただけの docx は not_memo_dump で落ちる。"""
    from docx import Document
    run_dir, _ = _run("word-01", tmp_path, with_expected=True)
    memo = (run_dir / "メモ.txt").read_text(encoding="utf-8").splitlines()
    doc = Document()
    for line in memo[:6]:      # 【…作成メモ】〜件名まで
        doc.add_paragraph(line)
    doc.add_paragraph("記")
    for line in memo[6:]:      # 日時:〜問い合わせ先: をラベル付きのまま並べる
        doc.add_paragraph(line)
    doc.add_paragraph("以上")
    doc.save(run_dir / "通知文.docx")
    items = {i["id"]: i for i in run_checker(ROOT, "word-01", run_dir)}
    assert items["header_fields"]["pass"] and items["body_content"]["pass"]
    assert items["ki_ijo_order"]["pass"] and items["phone"]["pass"]
    assert not items["not_memo_dump"]["pass"], items["not_memo_dump"]


def test_word01_memo_dump_without_labels_still_needs_numbered_items(tmp_path):
    """ラベルと「作成メモ」を消しても、記〜以上の間に番号付きの項目が無ければ落ちる。"""
    from docx import Document
    run_dir, _ = _run("word-01", tmp_path, with_expected=True)
    doc = Document()
    for line in ["中総第123号", "令和8年9月15日", "各自治会長 様", "みなと市役所 総務課長",
                 "令和8年度 地域防災訓練の実施について(案内)", "記",
                 "令和8年10月18日(日)に中央小学校グラウンドで実施します。持ち物は筆記用具です。",
                 "以上", "みなと市役所 総務課 防災係 電話 025-123-4567"]:
        doc.add_paragraph(line)
    doc.save(run_dir / "通知文.docx")
    items = {i["id"]: i for i in run_checker(ROOT, "word-01", run_dir)}
    assert items["ki_ijo_order"]["pass"]
    assert not items["not_memo_dump"]["pass"], items["not_memo_dump"]


def test_mail03_detects_wrong_slot_and_missing_recipient(tmp_path):
    run_dir, _ = _run("mail-03", tmp_path, with_expected=True)
    reply = run_dir / "reply.txt"
    good = reply.read_text(encoding="utf-8")
    reply.write_text(good.replace("9月17日(木) 11:00〜12:00", "9月15日(火) 10:00〜11:00"), encoding="utf-8")
    items = {i["id"]: i for i in run_checker(ROOT, "mail-03", run_dir)}
    assert not items["no_wrong_slot"]["pass"] and not items["all_correct_slots"]["pass"]
    assert items["to_all"]["pass"]
    reply.write_text(good.replace("takahashi@example.co.jp", ""), encoding="utf-8")
    items = {i["id"]: i for i in run_checker(ROOT, "mail-03", run_dir)}
    assert not items["to_all"]["pass"] and "takahashi" in items["to_all"]["detail"]


def test_mail03_extract_slots_accepts_fullwidth_and_variants():
    import importlib.util
    spec = importlib.util.spec_from_file_location("mail03_check", ROOT / "tasks" / "mail-03" / "check.py")
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    body = "候補① ９月１６日（水） １５：００～１６：００\n候補2 9月17日(木) 11:00-12:00\n"
    assert mod.extract_slots(body) == {(9, 16, 15, 0, 16, 0), (9, 17, 11, 0, 12, 0)}


def test_mail03_detects_wrong_slot_written_in_hour_format(tmp_path):
    """「10時〜11時」形式で書かれた都合の悪い候補も見逃さない。"""
    run_dir, _ = _run("mail-03", tmp_path, with_expected=True)
    reply = run_dir / "reply.txt"
    text = reply.read_text(encoding="utf-8")
    reply.write_text(text.replace("9月17日(木) 11:00〜12:00", "9月15日(火曜) 10時〜11時"), encoding="utf-8")
    items = {i["id"]: i for i in run_checker(ROOT, "mail-03", run_dir)}
    assert not items["no_wrong_slot"]["pass"] and "9月15日" in items["no_wrong_slot"]["detail"]
    assert not items["all_correct_slots"]["pass"]


def test_mail03_extract_slots_accepts_weekday_and_hour_variants():
    """「(水曜)」「15時〜16時」「11時00分〜12時00分」も候補として拾う。"""
    import importlib.util
    spec = importlib.util.spec_from_file_location("mail03_check", ROOT / "tasks" / "mail-03" / "check.py")
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    assert mod.extract_slots("9月16日(水曜) 15:00〜16:00") == {(9, 16, 15, 0, 16, 0)}
    assert mod.extract_slots("9月16日(水曜日) 15時〜16時") == {(9, 16, 15, 0, 16, 0)}
    assert mod.extract_slots("9月17日(木) 11時00分〜12時00分") == {(9, 17, 11, 0, 12, 0)}
    assert mod.extract_slots("9月16日(水) 15:00〜16時30分") == {(9, 16, 15, 0, 16, 30)}


def test_web02_detects_missing_row_wrong_weekday_and_bad_source(tmp_path):
    import openpyxl
    run_dir, _ = _run("web-02", tmp_path, with_expected=True)
    path = run_dir / "祝日2027.xlsx"
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    # 2行目(元日)の曜日を壊し、振替休日の行を消し、出典を別サイトにする
    ws.cell(row=2, column=2).value = "土"
    for row in range(2, ws.max_row + 1):
        if str(ws.cell(row=row, column=4).value or "").startswith("振替"):
            ws.delete_rows(row)
            break
    for row in range(1, ws.max_row + 1):
        for col in range(1, 5):
            v = ws.cell(row=row, column=col).value
            if isinstance(v, str) and "cao.go.jp" in v:
                ws.cell(row=row, column=col).value = "https://example.com/holidays"
    wb.save(path)
    items = {i["id"]: i for i in run_checker(ROOT, "web-02", run_dir)}
    assert items["xlsx_exists"]["pass"]
    assert not items["weekdays_consistent"]["pass"]
    assert not items["dates_match"]["pass"] and "2027-03-22" in items["dates_match"]["detail"]
    assert not items["substitute_marked"]["pass"]
    assert not items["source_is_cao"]["pass"]


def test_web02_accepts_string_dates_and_long_weekday(tmp_path):
    import openpyxl
    run_dir, _ = _run("web-02", tmp_path, with_expected=True)
    path = run_dir / "祝日2027.xlsx"
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    ws.cell(row=2, column=1).value = "2027/1/1"
    ws.cell(row=2, column=2).value = "金曜日"
    wb.save(path)
    items = {i["id"]: i for i in run_checker(ROOT, "web-02", run_dir)}
    assert items["dates_match"]["pass"] and items["weekdays_consistent"]["pass"]


def _replace_cell_containing(ws, marker: str, new_value: str) -> None:
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and marker in c.value:
                c.value = new_value
                return
    raise AssertionError(f"「{marker}」を含むセルが見つかりません")


def test_web02_rejects_spoofed_source_url(tmp_path):
    import openpyxl
    run_dir, _ = _run("web-02", tmp_path, with_expected=True)
    path = run_dir / "祝日2027.xlsx"
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    _replace_cell_containing(ws, "cao.go.jp", "出典: https://www.google.com/search?q=cao.go.jp")
    wb.save(path)
    items = {i["id"]: i for i in run_checker(ROOT, "web-02", run_dir)}
    assert not items["source_is_cao"]["pass"]

    wb = openpyxl.load_workbook(path)
    ws = wb.active
    _replace_cell_containing(ws, "google.com", "出典: https://www8.cao.go.jp/chosei/shukujitsu/gaiyou.html")
    wb.save(path)
    items = {i["id"]: i for i in run_checker(ROOT, "web-02", run_dir)}
    assert items["source_is_cao"]["pass"]


def test_web02_accepts_fullwidth_date_and_paren_weekday(tmp_path):
    import openpyxl
    run_dir, _ = _run("web-02", tmp_path, with_expected=True)
    path = run_dir / "祝日2027.xlsx"
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    ws.cell(row=2, column=1).value = "２０２７／１／１"
    ws.cell(row=2, column=2).value = "(金)"
    wb.save(path)
    items = {i["id"]: i for i in run_checker(ROOT, "web-02", run_dir)}
    assert items["dates_match"]["pass"] and items["weekdays_consistent"]["pass"]


def test_update01_detects_leftover_retiree_bad_numbering_and_modified_original(tmp_path):
    import openpyxl
    run_dir, _ = _run("update-01", tmp_path, with_expected=True)
    out = run_dir / "令和8年度_委員名簿.xlsx"
    wb = openpyxl.load_workbook(out)
    ws = wb.active
    ws.append([99, "公募委員", "小林 正雄", "北地区在住", "", "", "令和7年4月1日〜令和9年3月31日"])
    wb.save(out)
    src = run_dir / "令和7年度_委員名簿.xlsx"
    wb2 = openpyxl.load_workbook(src)
    wb2.active["A1"].value = "令和8年度 地域福祉推進協議会 委員名簿"
    wb2.save(src)
    items = {i["id"]: i for i in run_checker(ROOT, "update-01", run_dir)}
    assert items["file_exists"]["pass"]
    assert not items["original_untouched"]["pass"]
    assert not items["retired_absent"]["pass"] and "小林" in items["retired_absent"]["detail"]
    assert not items["numbering"]["pass"]
    assert items["new_members_correct"]["pass"] and items["changes_applied"]["pass"]


def test_update01_detects_wrong_term_and_old_phone(tmp_path):
    import openpyxl
    run_dir, _ = _run("update-01", tmp_path, with_expected=True)
    out = run_dir / "令和8年度_委員名簿.xlsx"
    wb = openpyxl.load_workbook(out)
    ws = wb.active
    for row in ws.iter_rows(min_row=1):
        cells = {str(c.value): c for c in row if c.value is not None}
        if "中村 恵子" in cells:
            row[6].value = "令和7年4月1日〜令和9年3月31日"
        if "渡辺 隆" in cells:
            row[5].value = "025-111-1001"
    wb.save(out)
    items = {i["id"]: i for i in run_checker(ROOT, "update-01", run_dir)}
    assert not items["new_members_correct"]["pass"] and "中村" in items["new_members_correct"]["detail"]
    assert not items["changes_applied"]["pass"] and "電話" in items["changes_applied"]["detail"]


def test_update01_detects_fullwidth_old_phone_leftover(tmp_path):
    import openpyxl
    run_dir, _ = _run("update-01", tmp_path, with_expected=True)
    out = run_dir / "令和8年度_委員名簿.xlsx"
    wb = openpyxl.load_workbook(out)
    ws = wb.active
    for row in ws.iter_rows(min_row=1):
        cells = {str(c.value): c for c in row if c.value is not None}
        if "渡辺 隆" in cells:
            row[5].value = "０２５-１１１-１００１"
    wb.save(out)
    items = {i["id"]: i for i in run_checker(ROOT, "update-01", run_dir)}
    assert not items["changes_applied"]["pass"] and "旧電話番号" in items["changes_applied"]["detail"]


def test_word01_accepts_parenthesized_and_fullwidth_numbering(tmp_path):
    """記書きの番号が（１）・全角数字・丸数字などの表記でも not_memo_dump は通る。"""
    import re
    from docx import Document
    run_dir, _ = _run("word-01", tmp_path, with_expected=True)
    path = run_dir / "通知文.docx"
    doc = Document(str(path))
    numbering = ["（１）", "(2)", "３", "4.", "5．", "⑥"]
    pattern = re.compile(r"^(\d) (.*)$")
    idx = 0
    for p in doc.paragraphs:
        m = pattern.match(p.text)
        if m:
            p.text = numbering[idx] + m.group(2)
            idx += 1
    assert idx == 6, f"番号付き項目が6件見つかりませんでした(見つかった数: {idx})"
    doc.save(path)
    items = {i["id"]: i for i in run_checker(ROOT, "word-01", run_dir)}
    assert items["not_memo_dump"]["pass"], items["not_memo_dump"]
    assert items["ki_ijo_order"]["pass"], items["ki_ijo_order"]
    assert items["body_content"]["pass"], items["body_content"]
